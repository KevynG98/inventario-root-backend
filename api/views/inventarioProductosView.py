from decimal import Decimal, InvalidOperation
from django.db.models import Q
from rest_framework.parsers import MultiPartParser, FormParser
from drf_yasg import openapi
from drf_yasg.utils import swagger_auto_schema
from rest_framework import status
from rest_framework.decorators import api_view, parser_classes
from rest_framework.response import Response
from openpyxl import load_workbook
import io
import time
from django.core.files.base import ContentFile

from api.utils.pagination import CustomPageNumberPagination
from ..models.inventarioProductoModel import InventarioProducto
from ..serializers.inventarioProductoSerializer import InventarioProductoSerializer


def _parse_decimal(value, default="0"):
    """
    Convierte valores de Excel a Decimal de forma tolerante.
    """
    if value is None:
        return Decimal(default)
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = str(value).strip().replace(",", "")
    if not text:
        return Decimal(default)
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return Decimal(default)


def _codigo_generado(existentes, contador):
    """
    Genera un código con prefijo IMP- que no exista en la base/local.
    """
    while True:
        codigo = f"IMP-{contador:05d}"
        contador += 1
        if codigo not in existentes:
            existentes.add(codigo)
            return codigo, contador


@swagger_auto_schema(
    method='get', 
    operation_summary="Listar productos activos con filtros", 
    tags=["inventario-productos"],
    manual_parameters=[
        openapi.Parameter('nombre', openapi.IN_QUERY, type=openapi.TYPE_STRING, description="Filtrar por nombre"),
        openapi.Parameter('min_price', openapi.IN_QUERY, type=openapi.TYPE_NUMBER, description="Precio mínimo"),
        openapi.Parameter('max_price', openapi.IN_QUERY, type=openapi.TYPE_NUMBER, description="Precio máximo"),
    ]
)
@api_view(['GET'])
def listar_productos(request):
    inventarios = InventarioProducto.objects.filter(is_active=True).order_by('id')

    # Filtros
    nombre = request.query_params.get('nombre')
    min_price = request.query_params.get('min_price')
    max_price = request.query_params.get('max_price')

    if nombre:
        inventarios = inventarios.filter(nombre__icontains=nombre)
    if min_price:
        try:
            inventarios = inventarios.filter(precio_stock__gte=min_price)
        except ValueError:
            pass
    if max_price:
        try:
            inventarios = inventarios.filter(precio_stock__lte=max_price)
        except ValueError:
            pass

    paginator = CustomPageNumberPagination()
    result_page = paginator.paginate_queryset(inventarios, request)
    serializer = InventarioProductoSerializer(result_page, many=True, context={'request': request})
    return paginator.get_paginated_response(serializer.data)


@swagger_auto_schema(method='post', request_body=InventarioProductoSerializer, operation_summary="Crear nuevo producto", tags=["inventario-productos"])
@api_view(['POST'])
def crear_producto(request):
    serializer = InventarioProductoSerializer(data=request.data, context={'request': request})
    serializer.is_valid(raise_exception=True)
    serializer.save()
    return Response(serializer.data, status=status.HTTP_201_CREATED)


@swagger_auto_schema(method='get', operation_summary="Obtener producto por ID", tags=["inventario-productos"])
@api_view(['GET'])
def obtener_producto(request, pk):
    try:
        inventario = InventarioProducto.objects.get(pk=pk)
    except InventarioProducto.DoesNotExist:
        return Response({"error": "Producto no encontrado"}, status=status.HTTP_404_NOT_FOUND)
    serializer = InventarioProductoSerializer(inventario, context={'request': request})
    return Response(serializer.data)


@swagger_auto_schema(method='put', request_body=InventarioProductoSerializer, operation_summary="Actualizar producto", tags=["inventario-productos"])
@api_view(['PUT'])
def actualizar_producto(request, pk):
    try:
        inventario = InventarioProducto.objects.get(pk=pk)
    except InventarioProducto.DoesNotExist:
        return Response({"error": "Producto no encontrado"}, status=status.HTTP_404_NOT_FOUND)

    serializer = InventarioProductoSerializer(inventario, data=request.data, partial=True, context={'request': request})
    if serializer.is_valid():
        serializer.save()
        return Response(serializer.data)
    print("❌ Error al actualizar producto:", serializer.errors)
    return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


@swagger_auto_schema(method='delete', operation_summary="Eliminar producto (soft delete)", tags=["inventario-productos"])
@api_view(['DELETE'])
def eliminar_producto(request, pk):
    try:
        inventario = InventarioProducto.objects.get(pk=pk)
    except InventarioProducto.DoesNotExist:
        return Response({"error": "Producto no encontrado"}, status=status.HTTP_404_NOT_FOUND)
    inventario.is_active = False
    inventario.save(update_fields=['is_active'])
    return Response({"mensaje": "Producto desactivado (soft delete)"}, status=status.HTTP_200_OK)


@swagger_auto_schema(
    method='get',
    operation_summary="Buscar productos",
    tags=["inventario-productos"],
    manual_parameters=[
        openapi.Parameter('q', openapi.IN_QUERY, description='Búsqueda libre por nombre o código de inventario', type=openapi.TYPE_STRING),
        openapi.Parameter('nombre', openapi.IN_QUERY, description='Filtrar por nombre (icontains)', type=openapi.TYPE_STRING),
        openapi.Parameter('codigo_inventario', openapi.IN_QUERY, description='Filtrar por código de inventario (icontains)', type=openapi.TYPE_STRING),
        openapi.Parameter('page', openapi.IN_QUERY, type=openapi.TYPE_INTEGER),
        openapi.Parameter('page_size', openapi.IN_QUERY, type=openapi.TYPE_INTEGER),
    ]
)
@api_view(['GET'])
def buscar_productos(request):
    queryset = InventarioProducto.objects.filter(is_active=True).order_by('nombre')

    q = request.query_params.get('q')
    nombre = request.query_params.get('nombre')
    codigo_inventario = request.query_params.get('codigo_inventario')

    filtros = Q()
    if q:
        ql = q.strip()
        if ql:
            filtros |= Q(nombre__icontains=ql)
            filtros |= Q(codigo_inventario__icontains=ql)
    if nombre:
        filtros &= Q(nombre__icontains=nombre)
    if codigo_inventario:
        filtros &= Q(codigo_inventario__icontains=codigo_inventario)

    if filtros:
        queryset = queryset.filter(filtros)

    paginator = CustomPageNumberPagination()
    result_page = paginator.paginate_queryset(queryset, request)
    serializer = InventarioProductoSerializer(result_page, many=True, context={'request': request})
    return paginator.get_paginated_response(serializer.data)


@swagger_auto_schema(
    method='post',
    operation_summary="Carga masiva de productos desde Excel (xlsx)",
    tags=["inventario-productos"],
    manual_parameters=[
        openapi.Parameter(
            'file',
            openapi.IN_FORM,
            description='Archivo Excel con columnas: Nombre, Imagen (incrustada)',
            type=openapi.TYPE_FILE,
            required=True
        )
    ],
    responses={
        200: openapi.Response(
            description="Resultado de la importación",
            examples={
                "application/json": {
                    "creados": 10,
                    "errores": [],
                    "resumen": "10 productos creados"
                }
            }
        )
    }
)
@api_view(['POST'])
@parser_classes([MultiPartParser, FormParser])
def carga_masiva_productos(request):
    """
    Recibe un Excel, extrae títulos e imágenes incrustadas, y crea los productos.
    """
    archivo = request.FILES.get('file')
    if not archivo:
        return Response({"error": "No se envió archivo con clave 'file'."}, status=status.HTTP_400_BAD_REQUEST)

    try:
        # Cargamos el libro SIN read_only para poder acceder a las imágenes
        libro = load_workbook(archivo, data_only=True)
        hoja = libro.active
    except Exception as exc:
        return Response({"error": "No se pudo leer el Excel", "detalle": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

    # Diccionario para mapear fila -> imagen
    imagenes_por_fila = {}
    
    # Intentar extraer imágenes incrustadas
    if hasattr(hoja, '_images'):
        for img in hoja._images:
            try:
                row = img.anchor._from.row + 1 
                img_data = img._data() 
                ext = "jpg"
                if hasattr(img, 'format') and img.format:
                    ext = img.format.lower()
                
                img_filename = f"import_{row}_{int(time.time()*1000)}.{ext}"
                imagenes_por_fila[row] = (img_filename, img_data)
            except Exception as e:
                print(f"Error extrayendo imagen: {e}")

    existentes = set(InventarioProducto.objects.values_list('codigo_inventario', flat=True))
    contador = len(existentes) + 1

    creados = []
    errores = []
    
    first_row = next(hoja.iter_rows(min_row=1, max_row=1, values_only=True), (None,))
    start_row = 1
    if first_row and str(first_row[0]).lower() in ['nombre', 'titulo', 'título', 'item']:
        start_row = 2

    for idx, fila in enumerate(hoja.iter_rows(min_row=start_row, values_only=True), start=start_row):
        if not fila or all(c is None for c in fila):
            continue
            
        nombre = fila[0]
        if not nombre or str(nombre).strip() == "":
            continue

        codigo, contador = _codigo_generado(existentes, contador)
        
        payload = {
            "nombre": str(nombre).strip(),
            "codigo_inventario": codigo,
            "precio_compre": 0,
            "precio_stock": 0,
            "is_active": True,
        }

        producto = InventarioProducto(**payload)
        
        if idx in imagenes_por_fila:
            img_filename, img_bytes = imagenes_por_fila[idx]
            producto.imagen.save(img_filename, ContentFile(img_bytes), save=False)

        try:
            producto.save()
            creados.append(producto.codigo_inventario)
        except Exception as e:
            errores.append({"fila": idx, "detalle": str(e)})

    resumen = f"{len(creados)} productos creados, {len(errores)} filas con errores"
    return Response(
        {
            "creados": len(creados),
            "errores": errores,
            "resumen": resumen,
        },
        status=status.HTTP_200_OK
    )